#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
import re
import os
from functools import reduce
from io import BytesIO
from timeit import default_timer as timer
from docx import Document
from docx.opc.pkgreader import _SerializedRelationships, _SerializedRelationship
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from docx.opc.oxml import parse_xml
from markdown import markdown
from PIL import Image
from common.token_utils import num_tokens_from_string

from common.constants import LLMType, MAXIMUM_PAGE_NUMBER
from api.db.services.llm_service import LLMBundle
from api.db.joint_services.tenant_model_service import (
    ensure_mineru_from_env,
    ensure_opendataloader_from_env,
    ensure_paddleocr_from_env,
    get_first_provider_model_name,
    get_model_config_from_provider_instance,
    get_tenant_default_model_by_type,
)
from rag.utils.file_utils import extract_embed_file, extract_links_from_pdf, extract_links_from_docx, extract_html
from deepdoc.parser import DocxParser, EpubParser, ExcelParser, HtmlParser, JsonParser, MarkdownElementExtractor, MarkdownParser, PdfParser, TxtParser
from deepdoc.parser.figure_parser import VisionFigureParser, vision_figure_parser_docx_wrapper_naive, vision_figure_parser_pdf_wrapper
from deepdoc.parser.pdf_parser import PlainParser, VisionParser
from deepdoc.parser.docling_parser import DoclingParser
from deepdoc.parser.tcadp_parser import TCADPParser
from common.float_utils import normalize_overlapped_percent
from common.parser_config_utils import normalize_layout_recognizer
from common.text_utils import normalize_arabic_presentation_forms
from rag.nlp import (
    concat_img,
    find_codec,
    naive_merge,
    naive_merge_with_images,
    naive_merge_docx,
    rag_tokenizer,
    tokenize_chunks,
    doc_tokenize_chunks_with_images,
    tokenize_table,
    append_context2table_image4pdf,
    tokenize_chunks_with_images,
)  # noqa: F401


def _is_short_header(text, max_tokens=50):
    """
    Check if text is a short markdown header.

    Args:
        text: The text to check
        max_tokens: Maximum tokens for a header to be considered "short"

    Returns:
        bool: True if text is a short markdown header, False otherwise
    """
    if not text or not text.strip():
        return False

    # Check if it matches markdown header pattern: 1-6 # followed by space
    if not re.match(r"^#{1,6}\s+", text.strip()):
        return False

    # Check if token count is below threshold
    return num_tokens_from_string(text) < max_tokens


def _normalize_section_text_for_rtl_presentation_forms(sections):
    if not sections:
        return sections

    normalized_sections = []
    for section in sections:
        if isinstance(section, tuple):
            if not section:
                normalized_sections.append(section)
                continue
            text = section[0]
            normalized_text = normalize_arabic_presentation_forms(text)
            normalized_sections.append((normalized_text, *section[1:]))
            continue
        if isinstance(section, list):
            if not section:
                normalized_sections.append(section)
                continue
            text = section[0]
            normalized_text = normalize_arabic_presentation_forms(text)
            normalized_sections.append([normalized_text, *section[1:]])
            continue
        normalized_sections.append(normalize_arabic_presentation_forms(section))

    return normalized_sections


def by_deepdoc(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    callback = callback
    binary = binary
    pdf_parser = pdf_cls() if pdf_cls else Pdf()
    sections, tables = pdf_parser(filename if not binary else binary, from_page=from_page, to_page=to_page, callback=callback)

    tables = vision_figure_parser_pdf_wrapper(
        tbls=tables,
        sections=sections,
        callback=callback,
        **kwargs,
    )
    return sections, tables, pdf_parser


def by_mineru(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    mineru_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not mineru_llm_name:
            try:
                mineru_llm_name = get_first_provider_model_name(tenant_id, "MinerU", LLMType.OCR) or ensure_mineru_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env mineru: {e}")

        if mineru_llm_name:
            try:
                ocr_model_config = get_model_config_from_provider_instance(tenant_id, LLMType.OCR, mineru_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl

                # Closes #14869: when the tenant has an IMAGE2TEXT model
                # configured, let the MinerU parser enrich image chunks with
                # VLM-generated semantic descriptions (parity with deepdoc's
                # VisionFigureParser). Best-effort — fall back silently if
                # no vision model is available.
                if "vision_model" not in kwargs:
                    try:
                        vision_model_config = get_tenant_default_model_by_type(tenant_id, LLMType.IMAGE2TEXT)
                        kwargs["vision_model"] = LLMBundle(tenant_id=tenant_id, model_config=vision_model_config, lang=lang)
                    except Exception as vlm_err:
                        logging.info(f"[MinerU] no IMAGE2TEXT model for tenant; skipping image VLM enhancement: {vlm_err}")

                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    lang=lang,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle MinerU ({mineru_llm_name}): {e}")

    if callback:
        callback(-1, "MinerU not found.")
    return None, None, None


def by_docling(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    pdf_parser = DoclingParser()
    parse_method = kwargs.get("parse_method", "raw")

    if not pdf_parser.check_installation():
        if callback:
            callback(-1, "Docling not found.")
        return None, None, pdf_parser

    sections, tables = pdf_parser.parse_pdf(
        filepath=filename,
        binary=binary,
        callback=callback,
        output_dir=os.environ.get("DOCLING_OUTPUT_DIR", ""),
        delete_output=bool(int(os.environ.get("DOCLING_DELETE_OUTPUT", 1))),
        docling_server_url=os.environ.get("DOCLING_SERVER_URL", ""),
        parse_method=parse_method,
    )
    return sections, tables, pdf_parser


def by_opendataloader(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    opendataloader_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    if tenant_id:
        if not opendataloader_llm_name:
            try:
                opendataloader_llm_name = get_first_provider_model_name(tenant_id, "OpenDataLoader", LLMType.OCR) or ensure_opendataloader_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env opendataloader: {e}")

        if opendataloader_llm_name:
            try:
                ocr_model_config = get_model_config_from_provider_instance(tenant_id, LLMType.OCR, opendataloader_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                parse_options = {k: kwargs[k] for k in ("hybrid", "image_output", "sanitize") if k in kwargs}
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **parse_options,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle OpenDataLoader ({opendataloader_llm_name}): {e}")

    if callback:
        callback(-1, "OpenDataLoader not found.")
    return None, None, None


def by_tcadp(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    tcadp_parser = TCADPParser()

    if not tcadp_parser.check_installation():
        callback(-1, "TCADP parser not available. Please check Tencent Cloud API configuration.")
        return None, None, tcadp_parser

    sections, tables = tcadp_parser.parse_pdf(filepath=filename, binary=binary, callback=callback, output_dir=os.environ.get("TCADP_OUTPUT_DIR", ""), file_type="PDF")
    return sections, tables, tcadp_parser


def by_paddleocr(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    paddleocr_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not paddleocr_llm_name:
            try:
                paddleocr_llm_name = get_first_provider_model_name(tenant_id, "PaddleOCR", LLMType.OCR) or ensure_paddleocr_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env paddleocr: {e}")

        if paddleocr_llm_name:
            try:
                ocr_model_config = get_model_config_from_provider_instance(tenant_id, LLMType.OCR, paddleocr_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle PaddleOCR ({paddleocr_llm_name}): {e}")

        return None, None, None

    if callback:
        callback(-1, "PaddleOCR not found.")
    return None, None, None


def by_somark(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    somark_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not somark_llm_name:
            try:
                from api.db.joint_services.tenant_model_service import ensure_somark_from_env

                somark_llm_name = ensure_somark_from_env(tenant_id)
            except Exception as e:
                logging.warning(f"fallback to env somark: {e}")

        if somark_llm_name:
            try:
                try:
                    ocr_model_config = get_model_config_from_provider_instance(tenant_id, LLMType.OCR, somark_llm_name)
                except Exception:
                    if "@" in somark_llm_name:
                        raise
                    from api.db.services.tenant_llm_service import TenantLLMService

                    ocr_model_config = TenantLLMService.get_model_config(tenant_id, LLMType.OCR.value, somark_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle SoMark ({somark_llm_name}): {e}")
                if callback:
                    callback(-1, f"Failed to parse pdf via SoMark ({somark_llm_name}): {e}")
                return None, None, None

    if callback:
        callback(-1, "SoMark not found.")
    return None, None, None


def by_plaintext(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, callback=None, **kwargs):
    layout_recognizer = (kwargs.get("layout_recognizer") or "").strip()
    if (not layout_recognizer) or (layout_recognizer == "Plain Text"):
        pdf_parser = PlainParser()
    else:
        tenant_id = kwargs.get("tenant_id")
        if not tenant_id:
            raise ValueError("tenant_id is required when using vision layout recognizer")
        vision_model_config = get_model_config_from_provider_instance(tenant_id, LLMType.IMAGE2TEXT, layout_recognizer)
        vision_model = LLMBundle(
            tenant_id,
            model_config=vision_model_config,
            lang=kwargs.get("lang", "Chinese"),
        )
        pdf_parser = VisionParser(vision_model=vision_model, **kwargs)

    sections, tables = pdf_parser(filename if not binary else binary, from_page=from_page, to_page=to_page, callback=callback)
    return sections, tables, pdf_parser


PARSERS = {
    "deepdoc": by_deepdoc,
    "mineru": by_mineru,
    "docling": by_docling,
    "opendataloader": by_opendataloader,
    "tcadp parser": by_tcadp,
    "paddleocr": by_paddleocr,
    "somark": by_somark,
    "plaintext": by_plaintext,  # default
}


class Docx(DocxParser):
    def __init__(self):
        pass

    def __clean(self, line):
        line = re.sub(r"\u3000", " ", line).strip()
        return line

    def __get_nearest_title(self, table_index, filename):
        """Get the hierarchical title structure before the table"""
        import re
        from docx.text.paragraph import Paragraph

        titles = []
        blocks = []

        # Get document name from filename parameter
        doc_name = re.sub(r"\.[a-zA-Z]+$", "", filename)
        if not doc_name:
            doc_name = "Untitled Document"

        # Collect all document blocks while maintaining document order
        try:
            # Iterate through all paragraphs and tables in document order
            for i, block in enumerate(self.doc._element.body):
                if block.tag.endswith("p"):  # Paragraph
                    p = Paragraph(block, self.doc)
                    blocks.append(("p", i, p))
                elif block.tag.endswith("tbl"):  # Table
                    blocks.append(("t", i, None))  # Table object will be retrieved later
        except Exception as e:
            logging.error(f"Error collecting blocks: {e}")
            return ""

        # Find the target table position
        target_table_pos = -1
        table_count = 0
        for i, (block_type, pos, _) in enumerate(blocks):
            if block_type == "t":
                if table_count == table_index:
                    target_table_pos = pos
                    break
                table_count += 1

        if target_table_pos == -1:
            return ""  # Target table not found

        # Find the nearest heading paragraph in reverse order
        nearest_title = None
        for i in range(len(blocks) - 1, -1, -1):
            block_type, pos, block = blocks[i]
            if pos >= target_table_pos:  # Skip blocks after the table
                continue

            if block_type != "p":
                continue

            if block.style and block.style.name and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                try:
                    level_match = re.search(r"(\d+)", block.style.name)
                    if level_match:
                        level = int(level_match.group(1))
                        if level <= 7:  # Support up to 7 heading levels
                            title_text = block.text.strip()
                            if title_text:  # Avoid empty titles
                                nearest_title = (level, title_text)
                                break
                except Exception as e:
                    logging.error(f"Error parsing heading level: {e}")

        if nearest_title:
            # Add current title
            titles.append(nearest_title)
            current_level = nearest_title[0]

            # Find all parent headings, allowing cross-level search
            while current_level > 1:
                found = False
                for i in range(len(blocks) - 1, -1, -1):
                    block_type, pos, block = blocks[i]
                    if pos >= target_table_pos:  # Skip blocks after the table
                        continue

                    if block_type != "p":
                        continue

                    if block.style and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                        try:
                            level_match = re.search(r"(\d+)", block.style.name)
                            if level_match:
                                level = int(level_match.group(1))
                                # Find any heading with a higher level
                                if level < current_level:
                                    title_text = block.text.strip()
                                    if title_text:  # Avoid empty titles
                                        titles.append((level, title_text))
                                        current_level = level
                                        found = True
                                        break
                        except Exception as e:
                            logging.error(f"Error parsing parent heading: {e}")

                if not found:  # Break if no parent heading is found
                    break

            # Sort by level (ascending, from highest to lowest)
            titles.sort(key=lambda x: x[0])
            # Organize titles (from highest to lowest)
            hierarchy = [doc_name] + [t[1] for t in titles]
            return " > ".join(hierarchy)

        return ""

    def __call__(self, filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER):
        self.doc = Document(filename) if not binary else Document(BytesIO(binary))
        pn = 0
        lines = []
        last_image = None
        table_idx = 0

        def flush_last_image():
            nonlocal last_image, lines
            if last_image is not None:
                lines.append({"text": "", "image": last_image, "table": None, "style": "Image"})
                last_image = None

        for block in self.doc._element.body:
            if pn > to_page:
                break

            if block.tag.endswith("p"):
                p = Paragraph(block, self.doc)

                if from_page <= pn < to_page:
                    text = p.text.strip()
                    style_name = p.style.name if p.style else ""

                    if text:
                        if style_name == "Caption":
                            former_image = None

                            if lines and lines[-1].get("image") and lines[-1].get("style") != "Caption":
                                former_image = lines[-1].get("image")
                                lines.pop()

                            elif last_image is not None:
                                former_image = last_image
                                last_image = None

                            lines.append(
                                {
                                    "text": self.__clean(text),
                                    "image": former_image if former_image else None,
                                    "table": None,
                                }
                            )

                        else:
                            flush_last_image()
                            lines.append(
                                {
                                    "text": self.__clean(text),
                                    "image": None,
                                    "table": None,
                                }
                            )

                            current_image = self.get_picture(self.doc, p)
                            if current_image is not None:
                                lines.append(
                                    {
                                        "text": "",
                                        "image": current_image,
                                        "table": None,
                                    }
                                )

                    else:
                        current_image = self.get_picture(self.doc, p)
                        if current_image is not None:
                            last_image = current_image

                for run in p.runs:
                    xml = run._element.xml
                    if "lastRenderedPageBreak" in xml:
                        pn += 1
                        continue
                    if "w:br" in xml and 'type="page"' in xml:
                        pn += 1

            elif block.tag.endswith("tbl"):
                if pn < from_page or pn > to_page:
                    table_idx += 1
                    continue

                flush_last_image()
                tb = DocxTable(block, self.doc)
                title = self.__get_nearest_title(table_idx, filename)
                html = "<table>"
                if title:
                    html += f"<caption>Table Location: {title}</caption>"
                for r in tb.rows:
                    html += "<tr>"
                    col_idx = 0
                    try:
                        while col_idx < len(r.cells):
                            span = 1
                            c = r.cells[col_idx]
                            for j in range(col_idx + 1, len(r.cells)):
                                if c.text == r.cells[j].text:
                                    span += 1
                                    col_idx = j
                                else:
                                    break
                            col_idx += 1
                            html += f"<td>{c.text}</td>" if span == 1 else f"<td colspan='{span}'>{c.text}</td>"
                    except Exception as e:
                        logging.warning(f"Error parsing table, ignore: {e}")
                    html += "</tr>"
                html += "</table>"
                lines.append({"text": "", "image": None, "table": html})
                table_idx += 1

        flush_last_image()
        new_line = [(line.get("text"), line.get("image"), line.get("table")) for line in lines]

        return new_line

    def to_markdown(self, filename=None, binary=None, inline_images: bool = True):
        """
        This function uses mammoth, licensed under the BSD 2-Clause License.
        """

        import base64
        import uuid

        import mammoth
        from markdownify import markdownify

        docx_file = BytesIO(binary) if binary else open(filename, "rb")

        def _convert_image_to_base64(image):
            try:
                with image.open() as image_file:
                    image_bytes = image_file.read()
                encoded = base64.b64encode(image_bytes).decode("utf-8")
                base64_url = f"data:{image.content_type};base64,{encoded}"

                alt_name = "image"
                alt_name = f"img_{uuid.uuid4().hex[:8]}"

                return {"src": base64_url, "alt": alt_name}
            except Exception as e:
                logging.warning(f"Failed to convert image to base64: {e}")
                return {"src": "", "alt": "image"}

        try:
            if inline_images:
                result = mammoth.convert_to_html(docx_file, convert_image=mammoth.images.img_element(_convert_image_to_base64))
            else:
                result = mammoth.convert_to_html(docx_file)

            html = result.value

            markdown_text = markdownify(html)
            return markdown_text

        finally:
            if not binary:
                docx_file.close()


class Pdf(PdfParser):
    def __init__(self):
        super().__init__()

    def __call__(self, filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, zoomin=3, callback=None, separate_tables_figures=False):
        start = timer()
        first_start = start
        callback(msg="OCR started")
        self.__images__(filename if not binary else binary, zoomin, from_page, to_page, callback)
        callback(msg="OCR finished ({:.2f}s)".format(timer() - start))
        logging.info("OCR({}~{}): {:.2f}s".format(from_page, to_page, timer() - start))

        start = timer()
        self._layouts_rec(zoomin)
        callback(0.63, "Layout analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._table_transformer_job(zoomin)
        callback(0.65, "Table analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._text_merge(zoomin=zoomin)
        callback(0.67, "Text merged ({:.2f}s)".format(timer() - start))

        if separate_tables_figures:
            tbls, figures = self._extract_table_figure(True, zoomin, True, True, True)
            self._concat_downward()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls, figures
        else:
            tbls = self._extract_table_figure(True, zoomin, True, True)
            self._naive_vertical_merge()
            self._concat_downward()
            # self._final_reading_order_merge()
            # self._filter_forpages()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls


# Maximum number of HTTP redirects followed when fetching a remote image
# referenced by a markdown document (each hop is SSRF-validated).
MAX_IMAGE_REDIRECTS = 5


class Markdown(MarkdownParser):
    def md_to_html(self, sections):
        if not sections:
            return []
        if isinstance(sections, type("")):
            text = sections
        elif isinstance(sections[0], type("")):
            text = sections[0]
        else:
            return []

        from bs4 import BeautifulSoup

        html_content = markdown(text)
        soup = BeautifulSoup(html_content, "html.parser")
        return soup

    def get_hyperlink_urls(self, soup):
        if soup:
            return set([a.get("href") for a in soup.find_all("a") if a.get("href")])
        return []

    def extract_image_urls_with_lines(self, text):
        md_img_re = re.compile(r"!\[[^\]]*\]\(([^)\s]+)")
        html_img_re = re.compile(r'src=["\\\']([^"\\\'>\\s]+)', re.IGNORECASE)
        urls = []
        seen = set()
        lines = text.splitlines()
        for idx, line in enumerate(lines):
            for url in md_img_re.findall(line):
                if (url, idx) not in seen:
                    urls.append({"url": url, "line": idx})
                    seen.add((url, idx))
            for url in html_img_re.findall(line):
                if (url, idx) not in seen:
                    urls.append({"url": url, "line": idx})
                    seen.add((url, idx))

        # cross-line
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "html.parser")
            newline_offsets = [m.start() for m in re.finditer(r"\n", text)] + [len(text)]
            for img_tag in soup.find_all("img"):
                src = img_tag.get("src")
                if not src:
                    continue

                tag_str = str(img_tag)
                pos = text.find(tag_str)
                if pos == -1:
                    # fallback
                    pos = max(text.find(src), 0)
                line_no = 0
                for i, off in enumerate(newline_offsets):
                    if pos <= off:
                        line_no = i
                        break
                if (src, line_no) not in seen:
                    urls.append({"url": src, "line": line_no})
                    seen.add((src, line_no))
        except Exception as e:
            logging.error("Failed to extract image urls: {}".format(e))
            pass

        return urls

    def load_images_from_urls(self, urls, cache=None):
        import requests
        from pathlib import Path
        from urllib.parse import urljoin

        from common.ssrf_guard import assert_url_is_safe, pin_dns

        cache = cache or {}
        images = []
        for url in urls:
            if url in cache:
                if cache[url]:
                    images.append(cache[url])
                continue
            img_obj = None
            try:
                if url.startswith(("http://", "https://")):
                    # SSRF guard: image references come from the (untrusted) uploaded
                    # document, so validate and DNS-pin every hop before connecting.
                    # Otherwise a markdown image like ![x](http://169.254.169.254/...)
                    # would make the server fetch internal services / cloud metadata.
                    # Redirects are followed manually so each hop is re-validated,
                    # mirroring common/data_source/rss_connector.py.
                    current_hostname, current_ip = assert_url_is_safe(url)
                    current_url = url
                    response = None
                    try:
                        for _ in range(MAX_IMAGE_REDIRECTS + 1):
                            # Release the previous hop before opening the next: with
                            # stream=True the connection isn't returned to the pool
                            # until the body is read or the response is closed.
                            if response is not None:
                                response.close()
                            with pin_dns(current_hostname, current_ip):
                                response = requests.get(current_url, stream=True, timeout=30, allow_redirects=False)
                            if response.status_code not in (301, 302, 303, 307, 308):
                                break
                            location = response.headers.get("Location")
                            if not location:
                                break
                            current_url = urljoin(current_url, location)
                            current_hostname, current_ip = assert_url_is_safe(current_url)
                        else:
                            raise ValueError(f"Exceeded {MAX_IMAGE_REDIRECTS} redirects fetching {url!r}")
                        if response.status_code == 200 and response.headers.get("Content-Type", "").startswith("image/"):
                            img_obj = Image.open(BytesIO(response.content)).convert("RGB")
                    finally:
                        # Always release the final/streamed response, including the
                        # non-image and redirect-cap paths where the body is unread.
                        if response is not None:
                            response.close()
                else:
                    local_path = Path(url)
                    if local_path.exists():
                        img_obj = Image.open(url).convert("RGB")
                    else:
                        logging.warning(f"Local image file not found: {url}")
            except Exception as e:
                logging.error(f"Failed to download/open image from {url}: {e}")
            cache[url] = img_obj
            if img_obj:
                images.append(img_obj)
        return images, cache

    def __call__(self, filename, binary=None, separate_tables=True, delimiter=None, return_section_images=False):
        """Parse markdown into text sections and optional standalone table chunks."""
        if binary:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(filename, "r") as f:
                txt = f.read()

        remainder, tables = self.extract_tables_and_remainder(f"{txt}\n", separate_tables=separate_tables)
        parsing_text = remainder if separate_tables else txt
        extractor = MarkdownElementExtractor(parsing_text)
        image_refs = self.extract_image_urls_with_lines(parsing_text)
        element_sections = extractor.extract_elements(delimiter, include_meta=True)

        sections = []
        section_images = []
        image_cache = {}
        for element in element_sections:
            content = element["content"]
            start_line = element["start_line"]
            end_line = element["end_line"]
            urls_in_section = [ref["url"] for ref in image_refs if start_line <= ref["line"] <= end_line]
            imgs = []
            if urls_in_section:
                imgs, image_cache = self.load_images_from_urls(urls_in_section, image_cache)
            combined_image = None
            if imgs:
                combined_image = reduce(concat_img, imgs) if len(imgs) > 1 else imgs[0]
            sections.append((content, ""))
            section_images.append(combined_image)

        tbls = []
        for table in tables:
            tbls.append(((None, markdown(table, extensions=["markdown.extensions.tables"])), ""))
        if return_section_images:
            return sections, tbls, section_images
        return sections, tbls


def load_from_xml_v2(baseURI, rels_item_xml):
    """
    Return |_SerializedRelationships| instance loaded with the
    relationships contained in *rels_item_xml*. Returns an empty
    collection if *rels_item_xml* is |None|.
    """
    srels = _SerializedRelationships()
    if rels_item_xml is not None:
        rels_elm = parse_xml(rels_item_xml)
        for rel_elm in rels_elm.Relationship_lst:
            if rel_elm.target_ref in ("../NULL", "NULL") or rel_elm.target_ref.startswith("#"):
                continue
            srels._srels.append(_SerializedRelationship(baseURI, rel_elm))
    return srels


def chunk(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, **kwargs):
    """
    使用通用（naive）策略解析并切割文档。

    支持 PDF、DOCX、Excel、TXT、Markdown、HTML、EPUB、JSON 等格式。
    文档会先按格式解析为连续文本段，再依据 delimiter 拆分，并按
    chunk_token_num 合并成最终分块；overlapped_percent 用于控制相邻分块重叠。
    """
    # 保存待递归处理的网页 URL，set 用于自动去重。
    urls = set()
    # 保存网页递归切割后产生的分块结果。
    url_res = []

    # 读取分块配置，并提前构建子分块所需的分隔符正则。
    # 未传语言时默认按中文处理。
    lang = lang or "Chinese"
    # 标记是否为英文，后续 Token 化时会选择对应语言规则。
    is_english = lang.lower() == "english"  # is_english(cks)
    # 从可变参数中读取解析/切割配置，没有配置时使用默认值。
    parser_config = kwargs.get("parser_config", {"chunk_token_num": 512, "delimiter": "\n!?。；！？", "layout_recognize": "DeepDOC", "analyze_hyperlink": True})

    # 将配置中的子块分隔符从转义字符串还原为实际文本。
    child_deli = (parser_config.get("children_delimiter") or "").encode("utf-8").decode("unicode_escape").encode("latin1").decode("utf-8")
    # 提取用反引号包裹的自定义子块分隔符。
    cust_child_deli = re.findall(r"`([^`]+)`", child_deli)
    # 将普通子分隔符转换成正则片段，并去掉自定义分隔符的反引号。
    child_deli = "|".join(re.sub(r"`([^`]+)`", "", child_deli))
    if cust_child_deli:
        # 长分隔符优先匹配，避免短分隔符抢先截断文本。
        cust_child_deli = sorted(set(cust_child_deli), key=lambda x: -len(x))
        # 对自定义分隔符进行正则转义，防止其中的特殊字符改变语义。
        cust_child_deli = "|".join(re.escape(t) for t in cust_child_deli if t)
        # 将自定义分隔符追加到普通子分隔符正则中。
        child_deli += cust_child_deli

    # 标记当前文件是否走 Markdown 专用合并逻辑。
    is_markdown = False
    # 读取表格上下文长度，负数统一按 0 处理。
    table_context_size = max(0, int(parser_config.get("table_context_size", 0) or 0))
    # 读取图片上下文长度，负数统一按 0 处理。
    image_context_size = max(0, int(parser_config.get("image_context_size", 0) or 0))

    # 初始化每个分块共享的文档元数据，标题会同时保存粗粒度和细粒度 Token。
    # 创建文档级元数据，并用文件名生成标题 Token。
    doc = {"docnm_kwd": filename, "title_tks": rag_tokenizer.tokenize(re.sub(r"\.[a-zA-Z]+$", "", filename))}
    # 为标题生成更细粒度的 Token，供检索排序使用。
    doc["title_sm_tks"] = rag_tokenizer.fine_grained_tokenize(doc["title_tks"])
    # 初始化最终分块结果列表。
    res = []
    # 保存 PDF 解析器实例，后续可能用于裁剪图片或读取大纲。
    pdf_parser = None
    # 保存 Markdown 章节对应的图片列表。
    section_images = None

    # 标记当前是否为根文档；递归处理附件和链接时会传 False。
    is_root = kwargs.get("is_root", True)
    # 保存嵌入附件递归切割得到的结果。
    embed_res = []
    if is_root:
        # 仅在根文档中提取附件，避免递归调用时重复扫描嵌入文件。
        # 初始化嵌入文件列表。
        embeds = []
        if binary is not None:
            # 从二进制文档中提取嵌入文件。
            embeds = extract_embed_file(binary)
        else:
            # 当前实现不支持只根据路径读取并提取嵌入文件。
            raise Exception("Embedding extraction from file path is not supported.")

        # 对每个嵌入文件递归执行相同切割流程，并汇总其分块结果。
        for embed_filename, embed_bytes in embeds:
            try:
                # 将嵌入文件作为非根文档递归切割，避免再次提取附件。
                sub_res = chunk(embed_filename, binary=embed_bytes, lang=lang, callback=callback, is_root=False, **kwargs) or []
                # 把嵌入文件的结果追加到附件结果集中。
                embed_res.extend(sub_res)
            except Exception as e:
                # 记录单个附件失败原因，不阻断主文档切割。
                error_msg = f"Failed to chunk embed {embed_filename}: {e}"
                # 将附件处理异常写入日志。
                logging.error(error_msg)
                if callback:
                    # 向调用方报告附件处理失败的进度消息。
                    callback(0.05, error_msg)
                # 跳过当前失败附件，继续处理下一个附件。
                continue

    if re.search(r"\.docx$", filename, re.IGNORECASE):
        # DOCX：解析段落、表格和图片，并保留它们之间的上下文关系。
        # 通知调用方 DOCX 解析开始。
        callback(0.1, "Start to parse.")
        if parser_config.get("analyze_hyperlink", False) and is_root:
            # 提取文档中的外部链接，链接页面会在主文档处理完成后递归切割。
            # 收集根 DOCX 中的网页链接。
            urls = extract_links_from_docx(binary)
            # 逐个处理 DOCX 中引用的网页。
            for index, url in enumerate(urls):
                # 下载网页二进制内容和元信息。
                html_bytes, metadata = extract_html(url)
                if not html_bytes:
                    # 下载失败时跳过当前网页。
                    continue
                try:
                    # 将网页作为非根文档递归切割。
                    sub_url_res = chunk(url, html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
                except Exception as e:
                    # URL 解析失败时记录日志，并用序号文件名重试。
                    logging.info(f"Failed to chunk url in registered file type {url}: {e}")
                    sub_url_res = chunk(f"{index}.html", html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
                # 汇总当前网页的分块结果。
                url_res.extend(sub_url_res)

        # 兼容 DOCX 中指向 NULL 或内部锚点的无效关系，避免读取压缩包时报错。
        # 替换 python-docx 的关系解析函数。
        _SerializedRelationships.load_from_xml = load_from_xml_v2

        # sections 中保留文本、图片和表格等结构化信息。
        # 读取 DOCX 的结构化内容。
        sections = Docx()(filename, binary)
        # 规范化从右向左语言的字符形式。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        # 按 Token 上限合并 DOCX 内容；images 记录包含图片的分块位置。
        # 将段落、表格和图片合并成最终的 DOCX 大块。
        chunks, images = naive_merge_docx(sections, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), table_context_size, image_context_size)

        # 如果配置了视觉模型，为图片分块补充语义描述。
        # 调用视觉解析器为图片补充文本描述。
        vision_figure_parser_docx_wrapper_naive(chunks=chunks, idx_lst=images, callback=callback, **kwargs)

        # 通知调用方 DOCX 文本解析完成。
        callback(0.8, "Finish parsing.")
        # 记录 DOCX Token 化开始时间。
        st = timer()

        # 将 DOCX 分块转成最终检索记录。
        res.extend(doc_tokenize_chunks_with_images(chunks, doc, is_english, child_delimiters_pattern=child_deli, language=lang))
        # 输出 DOCX 合并和 Token 化耗时。
        logging.info("naive_merge({}): {}".format(filename, timer() - st))
        # 合并嵌入附件的结果。
        res.extend(embed_res)
        # 合并外链网页的结果。
        res.extend(url_res)
        # DOCX 分支无需执行后面的通用流程。
        return res

    elif re.search(r"\.pdf$", filename, re.IGNORECASE):
        # PDF：根据 layout_recognize 选择 DeepDOC、MinerU、纯文本等解析器。
        # 标准化布局解析器配置并提取模型名。
        layout_recognizer, parser_model_name = normalize_layout_recognizer(parser_config.get("layout_recognize", "DeepDOC"))
        # 读取 OpenDataLoader 的可选模型名。
        opendataloader_llm_name = kwargs.pop("opendataloader_llm_name", None)
        if layout_recognizer == "OpenDataLoader" and parser_model_name:
            # OpenDataLoader 明确指定模型时覆盖外部传入值。
            opendataloader_llm_name = parser_model_name

        if parser_config.get("analyze_hyperlink", False) and is_root:
            # 收集根 PDF 中的网页链接。
            urls = extract_links_from_pdf(binary)

        if isinstance(layout_recognizer, bool):
            # 兼容布尔配置：True 表示 DeepDOC，False 表示纯文本。
            layout_recognizer = "DeepDOC" if layout_recognizer else "PlainText"

        # 转成小写名称，用于查找解析器。
        name = layout_recognizer.strip().lower()
        # 未找到指定解析器时回退到纯文本解析。
        parser = PARSERS.get(name, by_plaintext)
        # 通知调用方 PDF 解析开始。
        callback(0.1, "Start to parse.")

        # 执行 PDF 解析，获得正文、表格和解析器实例。
        sections, tables, pdf_parser = parser(
            filename=filename,
            binary=binary,
            from_page=from_page,
            to_page=to_page,
            lang=lang,
            callback=callback,
            layout_recognizer=layout_recognizer,
            mineru_llm_name=parser_model_name,
            paddleocr_llm_name=parser_model_name,
            opendataloader_llm_name=opendataloader_llm_name,
            somark_llm_name=parser_model_name,
            **kwargs,
        )
        # 统一正文中的阿拉伯语展示字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        if not sections and not tables:
            # 正文和表格都为空时直接返回。
            return []

        if table_context_size or image_context_size:
            # 为表格和图片补充附近正文，使其单独检索时仍保留必要上下文。
            # 将附近文本上下文附加到媒体块。
            tables = append_context2table_image4pdf(sections, tables, image_context_size)

        if name in ["tcadp", "docling", "mineru", "paddleocr", "opendataloader", "somark"]:
            # 这些解析器可能已经完成结构化切割。
            if int(parser_config.get("chunk_token_num", 0)) <= 0:
                # 使用 0 表示后续不强制按 Token 再切一次。
                parser_config["chunk_token_num"] = 0

        # 将 PDF 表格转换为可检索结果。
        res = tokenize_table(tables, doc, is_english, language=lang)
        # 通知调用方 PDF 解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(csv|xlsx?)$", filename, re.IGNORECASE):
        # 表格文件可使用 TCADP，也可回退到默认 ExcelParser。
        # 通知调用方表格解析开始。
        callback(0.1, "Start to parse.")

        # 根据配置判断电子表格是否使用 TCADP 解析器。
        layout_recognizer = parser_config.get("layout_recognize", "DeepDOC")
        if layout_recognizer == "TCADP Parser":
            # 读取 TCADP 的表格输出类型配置。
            table_result_type = parser_config.get("table_result_type", "1")
            # 读取 TCADP 的 Markdown 图片返回类型配置。
            markdown_image_response_type = parser_config.get("markdown_image_response_type", "1")
            # 创建 TCADP 解析器实例。
            tcadp_parser = TCADPParser(table_result_type=table_result_type, markdown_image_response_type=markdown_image_response_type)
            if not tcadp_parser.check_installation():
                # 依赖不可用时报告错误并返回当前已有结果。
                callback(-1, "TCADP parser not available. Please check Tencent Cloud API configuration.")
                return res

            # TCADP 需要显式区分 XLS/XLSX 与 CSV 输入类型。
            # 根据扩展名确定远程解析接口的文件类型。
            file_type = "XLSX" if re.search(r"\.xlsx?$", filename, re.IGNORECASE) else "CSV"

            # 调用 TCADP 解析表格，并取得正文和表格结果。
            sections, tables = tcadp_parser.parse_pdf(filepath=filename, binary=binary, callback=callback, output_dir=os.environ.get("TCADP_OUTPUT_DIR", ""), file_type=file_type)
            # 规范化解析结果中的 RTL 字符。
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)
            # TCADP 已经按表格结构处理，关闭后续 Token 二次切分。
            parser_config["chunk_token_num"] = 0
            # 将 TCADP 表格转换为检索记录。
            res = tokenize_table(tables, doc, is_english, language=lang)
            # 通知调用方表格解析完成。
            callback(0.8, "Finish parsing.")
        else:
            # 默认解析器可输出 HTML 表格，也可按普通行文本输出。
            # 创建默认 Excel 解析器。
            excel_parser = ExcelParser()
            if parser_config.get("html4excel"):
                # 将工作表转换为 HTML 文本，适合保留表格结构。
                sections = [(_, "") for _ in excel_parser.html(binary, 12) if _]
                # HTML 表格已经完成结构化处理，关闭二次切分。
                parser_config["chunk_token_num"] = 0
            else:
                # 按普通文本行读取 Excel 内容。
                sections = [(_, "") for _ in excel_parser(binary) if _]
            # 规范化 Excel 文本中的 RTL 字符。
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)

    elif re.search(r"\.(txt|py|js|java|c|cpp|h|php|go|ts|sh|cs|kt|sql)$", filename, re.IGNORECASE):
        # 纯文本和源码文件直接按配置的分隔符及 Token 上限预切分。
        # 通知调用方文本解析开始。
        callback(0.1, "Start to parse.")
        # 使用文本解析器按分隔符预切出 sections。
        sections = TxtParser()(filename, binary, parser_config.get("chunk_token_num", 128), parser_config.get("delimiter", "\n!?;。；！？"))
        # 规范化文本中的 RTL 字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        # 输出调试分隔线。
        print("\n", "-" * 150, "\n")
        # 输出预切分结果，便于调试文本解析。
        print(sections)
        # 输出调试分隔线。
        print("\n", "-" * 150, "\n")
        # 通知调用方文本解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(md|markdown|mdx)$", filename, re.IGNORECASE):
        # Markdown 单独保留章节、表格和图片，后续按标题语义进行合并。
        # 通知调用方 Markdown 解析开始。
        callback(0.1, "Start to parse.")
        # 创建 Markdown 解析器，并传入目标 Token 大小。
        markdown_parser = Markdown(int(parser_config.get("chunk_token_num", 128)))
        # 提取 Markdown 正文、表格和章节图片。
        sections, tables, section_images = markdown_parser(
            filename,
            binary,
            separate_tables=True,
            delimiter=parser_config.get("delimiter", "\n!?;。；！？"),
            return_section_images=True,
        )
        # 规范化 Markdown 文本中的 RTL 字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        # 标记后续使用 Markdown 专用合并算法。
        is_markdown = True

        try:
            # 获取租户默认的图片转文字模型配置。
            vision_model_config = get_tenant_default_model_by_type(kwargs["tenant_id"], LLMType.IMAGE2TEXT)
            # 创建图片语义增强模型。
            vision_model = LLMBundle(kwargs["tenant_id"], vision_model_config)
            # 报告已检测到视觉模型。
            callback(0.2, "Visual model detected. Attempting to enhance figure extraction...")
        except Exception as e:
            # 视觉模型不可用时记录原因并禁用图片增强。
            logging.warning(f"Failed to detect figure extraction: {e}")
            vision_model = None

        if vision_model:
            # 逐章节处理图片，并把视觉模型生成的描述追加到对应章节文本中。
            # 按章节索引遍历正文。
            for idx, (section_text, _) in enumerate(sections):
                # 初始化当前章节的图片列表。
                images = []
                if section_images and len(section_images) > idx and section_images[idx] is not None:
                    # 收集当前章节对应的图片。
                    images.append(section_images[idx])

                if images and len(images) > 0:
                    # 同一章节存在多张图片时先拼接，再统一交给视觉模型解析。
                    # 将章节中的多张图片合成为一张待识别图片。
                    combined_image = reduce(concat_img, images) if len(images) > 1 else images[0]
                    if section_images:
                        # 更新已有图片数组中的当前章节图片。
                        section_images[idx] = combined_image
                    else:
                        # 没有图片数组时按章节数创建占位数组。
                        section_images = [None] * len(sections)
                        # 将合成图片放入当前章节位置。
                        section_images[idx] = combined_image
                    # 创建 Markdown 图片视觉解析器。
                    markdown_vision_parser = VisionFigureParser(vision_model=vision_model, figures_data=[((combined_image, ["markdown image"]), [(0, 0, 0, 0, 0)])], **kwargs)
                    # 调用视觉模型生成图片描述。
                    boosted_figures = markdown_vision_parser(callback=callback)
                    # 把图片描述追加到章节正文末尾。
                    sections[idx] = (section_text + "\n\n" + "\n\n".join([fig[0][1] for fig in boosted_figures]), sections[idx][1])

        else:
            # 没有视觉模型时跳过图片语义增强。
            logging.warning("No visual model detected. Skipping figure parsing enhancement.")

        if parser_config.get("hyperlink_urls", False) and is_root:
            # 仅根 Markdown 文档提取其中的超链接。
            for idx, (section_text, _) in enumerate(sections):
                # 将当前 Markdown 章节转换为 HTML。
                soup = markdown_parser.md_to_html(section_text)
                # 从 HTML 中提取网页 URL。
                hyperlink_urls = markdown_parser.get_hyperlink_urls(soup)
                # 合并 URL 并自动去重。
                urls.update(hyperlink_urls)
        # 先将 Markdown 表格转为检索记录。
        res = tokenize_table(tables, doc, is_english, language=lang)
        # 通知调用方 Markdown 解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(htm|html)$", filename, re.IGNORECASE):
        # HTML 解析器会过滤标签，并尽量保持块级元素的文本边界。
        # 通知调用方 HTML 解析开始。
        callback(0.1, "Start to parse.")
        # 读取 HTML 使用的 Token 上限。
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        # 解析 HTML 并按块返回正文。
        sections = HtmlParser()(filename, binary, chunk_token_num)
        # 过滤空文本，并补充空位置信息。
        sections = [(_, "") for _ in sections if _]
        # 规范化 HTML 文本中的 RTL 字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        # 通知调用方 HTML 解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.epub$", filename, re.IGNORECASE):
        # EPUB 按内部章节顺序提取正文。
        # 通知调用方 EPUB 解析开始。
        callback(0.1, "Start to parse.")
        # 读取 EPUB 使用的 Token 上限。
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        # 解析 EPUB 内部章节。
        sections = EpubParser()(filename, binary, chunk_token_num)
        # 过滤空章节，并补充空位置信息。
        sections = [(_, "") for _ in sections if _]
        # 规范化 EPUB 文本中的 RTL 字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        # 通知调用方 EPUB 解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(json|jsonl|ldjson)$", filename, re.IGNORECASE):
        # JSON 系列格式按数据结构拆分，避免简单字符串截断破坏对象边界。
        # 通知调用方 JSON 解析开始。
        callback(0.1, "Start to parse.")
        # 读取 JSON 使用的 Token 上限。
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        # 根据 JSON 数据结构进行解析和切分。
        sections = JsonParser(chunk_token_num)(binary)
        # 过滤空结果，并补充空位置信息。
        sections = [(_, "") for _ in sections if _]
        # 规范化 JSON 文本中的 RTL 字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        # 通知调用方 JSON 解析完成。
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.doc$", filename, re.IGNORECASE):
        # 旧版 DOC 依赖 Tika 提取纯文本，再进入通用合并流程。
        # 通知调用方旧版 DOC 解析开始。
        callback(0.1, "Start to parse.")

        try:
            # 延迟导入 Tika，避免未安装时影响其他格式。
            from tika import parser as tika_parser
        except Exception as e:
            # Tika 不可用时报告不支持 DOC 解析。
            callback(0.8, f"tika not available: {e}. Unsupported .doc parsing.")
            # 将依赖缺失原因写入日志。
            logging.warning(f"tika not available: {e}. Unsupported .doc parsing for {filename}.")
            # 旧版 DOC 无法处理，返回空结果。
            return []

        # 将二进制内容包装成 Tika 可读取的文件流。
        binary = BytesIO(binary)
        # 使用 Tika 从文件流提取文本。
        doc_parsed = tika_parser.from_buffer(binary)
        if doc_parsed.get("content", None) is not None:
            # 按换行拆分 Tika 提取出的正文。
            sections = doc_parsed["content"].split("\n")
            # 删除空行并补充空位置信息。
            sections = [(_, "") for _ in sections if _]
            # 规范化 DOC 文本中的 RTL 字符。
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)
            # 通知调用方 DOC 解析完成。
            callback(0.8, "Finish parsing.")
        else:
            # 构造空正文错误信息。
            error_msg = f"tika.parser got empty content from {filename}."
            # 向调用方报告解析结果为空。
            callback(0.8, error_msg)
            # 记录空正文警告。
            logging.warning(error_msg)
            # 没有正文时返回空结果。
            return []
    else:
        # 当前文件扩展名不在支持列表中。
        raise NotImplementedError("file type not supported yet(pdf, xlsx, doc, docx, txt supported)")

    # 记录通用分块阶段的开始时间。
    st = timer()
    # 规范化分块重叠比例，防止传入异常值。
    overlapped_percent = normalize_overlapped_percent(parser_config.get("overlapped_percent", 0))

    if is_markdown:
        # Markdown 使用专用合并逻辑：短标题必须与其后的正文保持在同一分块中。
        # 保存合并后的 Markdown 文本块。
        merged_chunks = []
        # 保存每个文本块对应的合并图片。
        merged_images = []
        # 获取 Markdown 分块 Token 上限。
        chunk_limit = max(0, int(parser_config.get("chunk_token_num", 128)))

        # 初始化当前正在构建的文本块。
        current_text = ""
        # 初始化当前文本块的 Token 计数。
        current_tokens = 0
        # 初始化当前文本块的图片。
        current_image = None

        # 按章节顺序逐个合并 Markdown sections。
        for idx, sec in enumerate(sections):
            # 兼容带位置信息的 tuple 和纯字符串 section。
            text = sec[0] if isinstance(sec, tuple) else sec
            # 统计当前章节的 Token 数。
            sec_tokens = num_tokens_from_string(text)
            # 按章节索引取出对应图片。
            sec_image = section_images[idx] if section_images and idx < len(section_images) else None

            # 当前块仅为短标题时不立即封口，强制与下一节正文合并。
            if current_text and not _is_short_header(current_text) and current_tokens + sec_tokens > chunk_limit:
                # 当前块达到上限时加入已完成列表。
                merged_chunks.append(current_text)
                # 保存当前块对应的图片。
                merged_images.append(current_image)
                # 默认新块不带重叠文本。
                overlap_part = ""
                if overlapped_percent > 0:
                    # 将上一块末尾的一部分带入新块，减少跨块语义断裂。
                    # 按字符长度计算需要重叠的尾部长度。
                    overlap_len = int(len(current_text) * overlapped_percent / 100)
                    if overlap_len > 0:
                        # 截取上一块末尾作为新块前缀。
                        overlap_part = current_text[-overlap_len:]
                # 用重叠文本初始化下一个块。
                current_text = overlap_part
                # 重新计算新块当前的 Token 数。
                current_tokens = num_tokens_from_string(current_text)
                # 只有存在重叠文本时才保留上一块图片。
                current_image = current_image if overlap_part else None

            if current_text:
                # 当前块已有内容时，用换行追加新章节。
                current_text += "\n" + text
            else:
                # 当前块为空时直接使用新章节作为起始内容。
                current_text = text
            # 累加当前章节的 Token 数。
            current_tokens += sec_tokens

            if sec_image:
                # 将当前章节图片并入当前文本块。
                current_image = concat_img(current_image, sec_image) if current_image else sec_image

        if current_text:
            # 循环结束后，把最后一个未封口的块加入结果。
            merged_chunks.append(current_text)
            # 保存最后一个块对应的图片。
            merged_images.append(current_image)

        # 将专用合并结果作为后续 Token 化输入。
        chunks = merged_chunks
        # 判断 Markdown 结果中是否至少存在一张图片。
        has_images = merged_images and any(img is not None for img in merged_images)

        if has_images:
            # 图片与正文使用相同的分块索引，Token 化时一并写入结果。
            # 生成带图片信息的检索分块。
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, merged_images, child_delimiters_pattern=child_deli, language=lang))
        else:
            # 没有图片时生成普通文本检索分块。
            res.extend(tokenize_chunks(chunks, doc, is_english, pdf_parser, child_delimiters_pattern=child_deli, language=lang))
    else:
        # 非 Markdown 内容统一交给 naive_merge，按分隔符、Token 上限和重叠比例合并。
        # 如果图片列表全部为空，则按无图片文本处理。
        if section_images:
            if all(image is None for image in section_images):
                section_images = None

        if section_images:
            # 带图片的 sections 使用图片感知的合并函数。
            chunks, images = naive_merge_with_images(sections, section_images, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), overlapped_percent)
            # 将文本块和图片一起转换为检索记录。
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images, child_delimiters_pattern=child_deli, language=lang))
        else:
            # 不带图片的 sections 使用普通 naive_merge 合并。
            chunks = naive_merge(sections, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), overlapped_percent)

            # 将普通文本块转换为最终检索记录。
            res.extend(tokenize_chunks(chunks, doc, is_english, pdf_parser, child_delimiters_pattern=child_deli, language=lang))

    if urls and parser_config.get("analyze_hyperlink", False) and is_root:
        # 下载并递归切割正文中发现的网页链接，最终与主文档分块合并。
        # 按顺序处理去重后的网页 URL。
        for index, url in enumerate(urls):
            # 下载网页正文。
            html_bytes, metadata = extract_html(url)
            if not html_bytes:
                # 网页下载失败时跳过当前 URL。
                continue
            try:
                # 使用原始 URL 作为递归文档名切割网页。
                sub_url_res = chunk(url, html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
            except Exception as e:
                # 原始 URL 解析失败时记录日志。
                logging.info(f"Failed to chunk url in registered file type {url}: {e}")
                # 使用序号 HTML 文件名再次执行递归切割。
                sub_url_res = chunk(f"{index}.html", html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
            # 合并网页递归切割结果。
            url_res.extend(sub_url_res)

    # 输出当前文档从合并到结果生成的耗时。
    logging.info("naive_merge({}): {}".format(filename, timer() - st))

    if embed_res:
        # 如果存在嵌入附件，则追加附件分块。
        res.extend(embed_res)
    if url_res:
        # 如果存在外链网页，则追加网页分块。
        res.extend(url_res)
    # if table_context_size or image_context_size:
    #    attach_media_context(res, table_context_size, image_context_size)

    # 将 PDF 大纲临时挂到首个分块，随后由 task_executor.py 提取并持久化为文档元数据。
    if res and pdf_parser and getattr(pdf_parser, "outlines", None):
        # 将标题和层级转换成简单字典，便于后续序列化。
        res[0]["__outline__"] = [{"title": title, "depth": depth} for title, depth, *_ in pdf_parser.outlines]

    # 返回当前文档、附件和外链网页的全部分块。
    return res


if __name__ == "__main__":
    import sys

    def dummy(prog=None, msg=""):
        pass

    chunk(sys.argv[1], from_page=0, to_page=10, callback=dummy)
