#!/usr/bin/env python3
"""RAGFlow table.py 的独立表格解析、千问向量化和 Chroma 检索示例。

核心规则与 table.py 一致：第一行是字段名，每一行数据直接成为一个 chunk，
不会按照 chunk_size 再切分或合并。每行会转换成如下可检索文本：

    - 商品名称: 无线键盘
    - 分类: 办公设备
    - 描述: 适合日常办公

CSV 和制表符 TXT/TSV 仅使用 Python 标准库；XLSX 需要 openpyxl。
本 demo 不依赖 RAGFlow 服务、Pandas、知识库配置或 Elasticsearch。

运行示例：
    python table_demo.py --rows-only
    export DASHSCOPE_API_KEY='your-api-key'
    python table_demo.py --query '哪些办公设备库存不足？'
    python table_demo.py other.csv --rows-only

不传文件时默认读取脚本同目录下的 table_test.csv。向量默认保存到同目录的
table_chroma_db/，collection 名称为 table-demo。
"""

from __future__ import annotations

import argparse
import csv
import getpass
import hashlib
import io
import json
import math
import os
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen


DEFAULT_BASE_URL = "https://ws-leyp0n1kyjsllxss.cn-beijing.maas.aliyuncs.com/compatible-mode/v1"
DEFAULT_MODEL = "text-embedding-v2"
# 默认测试文件路径直接写在代码中；命令行传入其他文件时可以覆盖。
DEFAULT_DOCUMENT = Path(__file__).resolve().parent / "table_test.csv"
DEFAULT_CHROMA_PATH = Path(__file__).resolve().parent / "table_chroma_db"
DEFAULT_COLLECTION_NAME = "table-demo"
DEFAULT_QUERY = "哪些办公设备库存不足？"
DEFAULT_TOP_K = 3
DEFAULT_BATCH_SIZE = 4
DEFAULT_TIMEOUT = 60.0

# 可以在这里配置列的用途。未配置的列默认为 both。
# indexing/vectorize：只进入向量文本；metadata：只保存元数据；both：两者都进入。
COLUMN_ROLES: dict[str, str] = {}
INDEXING_ROLES = {"indexing", "vectorize", "both"}
METADATA_ROLES = {"metadata", "both"}
RESERVED_COLUMNS = {"id", "_id", "index", "idx"}
BOOL_TRUE = {"true", "yes", "是", "1", "y"}
BOOL_FALSE = {"false", "no", "否", "0", "n"}


class DemoError(RuntimeError):
    """表格解析、向量请求或 Chroma 操作失败。"""


@dataclass(frozen=True)
class TableRecord:
    """表格中的一行数据及其源文件行号。"""

    row_number: int
    values: dict[str, object]


@dataclass(frozen=True)
class ParsedTable:
    """解析后的表头、字段类型、有效记录和失败行号。"""

    name: str
    headers: list[str]
    column_types: dict[str, str]
    records: list[TableRecord]
    failed_rows: list[int]


def _decode(data: bytes) -> str:
    """依次尝试常见中英文编码。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize_headers(values: list[object]) -> list[str]:
    """清理表头，为缺失字段生成 Column_N，并拒绝重复字段名。"""
    headers = [str(value).strip() if value is not None and str(value).strip() else f"Column_{index}" for index, value in enumerate(values, 1)]
    duplicates = sorted({header for header in headers if headers.count(header) > 1})
    if duplicates:
        raise DemoError(f"存在重复字段名：{', '.join(duplicates)}")
    return headers


def _read_delimited(path: Path, delimiter: str) -> tuple[list[str], list[tuple[int, list[object]]], list[int]]:
    """读取 CSV/TSV，并跳过列数不匹配的行。CSV 引号由 csv.reader 处理。"""
    reader = csv.reader(io.StringIO(_decode(path.read_bytes())), delimiter=delimiter)
    try:
        headers = _normalize_headers(next(reader))
    except StopIteration as exc:
        raise DemoError("表格文件为空") from exc
    rows: list[tuple[int, list[object]]] = []
    failed_rows: list[int] = []
    for row_number, row in enumerate(reader, start=2):
        if not row or not any(str(value).strip() for value in row):
            continue
        if len(row) != len(headers):
            failed_rows.append(row_number)
            continue
        rows.append((row_number, row))
    return headers, rows, failed_rows


def _read_xlsx(path: Path) -> tuple[list[str], list[tuple[int, list[object]]], list[int]]:
    """读取 XLSX 的第一个工作表；旧版 XLS 请先转换为 XLSX。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DemoError("解析 XLSX 需要安装 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = _normalize_headers(list(next(iterator)))
        except StopIteration as exc:
            raise DemoError("Excel 文件为空") from exc
        rows: list[tuple[int, list[object]]] = []
        failed_rows: list[int] = []
        for row_number, row in enumerate(iterator, start=2):
            values = list(row)
            if not any(value is not None and str(value).strip() for value in values):
                continue
            # openpyxl 可能省略末尾空单元格，这里补齐到表头长度。
            values.extend([None] * (len(headers) - len(values)))
            if len(values) != len(headers):
                failed_rows.append(row_number)
                continue
            rows.append((row_number, values))
        return headers, rows, failed_rows
    finally:
        workbook.close()


def _parse_datetime(value: str) -> str | None:
    """识别常见 ISO 日期和日期时间，统一输出 ISO 字符串。"""
    text = value.strip().replace("Z", "+00:00")
    for candidate in (text, text.replace("/", "-")):
        try:
            return datetime.fromisoformat(candidate).isoformat(sep=" ")
        except ValueError:
            continue
    return None


def infer_column_type(values: list[object]) -> str:
    """按照 bool、int、float、datetime、text 的顺序推断整列类型。"""
    texts = [str(value).strip() for value in values if value is not None and str(value).strip()]
    if not texts:
        return "text"
    lowered = [value.lower() for value in texts]
    if all(value in BOOL_TRUE | BOOL_FALSE for value in lowered):
        return "bool"
    # 以 0 开头的编号、电话等保持文本，避免丢失前导零。
    if all(re.fullmatch(r"[+-]?(?:0|[1-9][0-9]*)", value) and not re.fullmatch(r"[+-]?0[0-9]+", value) for value in texts):
        return "int"
    if all(re.fullmatch(r"[+-]?(?:[0-9]+\.[0-9]+|[1-9][0-9]*|0)", value) and not re.fullmatch(r"[+-]?0[0-9]+", value) for value in texts):
        return "float"
    if all(_parse_datetime(value) is not None for value in texts):
        return "datetime"
    return "text"


def _convert_value(value: object, value_type: str) -> object | None:
    """将单元格转换成适合 Chroma metadata 保存的标量。"""
    if value is None or not str(value).strip():
        return None
    text = str(value).strip()
    try:
        if value_type == "bool":
            return text.lower() in BOOL_TRUE
        if value_type == "int":
            return int(text)
        if value_type == "float":
            return float(text)
        if value_type == "datetime":
            return _parse_datetime(text) or text
    except ValueError:
        return text
    return text


def read_table(path: Path | None = None) -> ParsedTable:
    """读取表格并推断字段类型；CSV/TXT/XLSX 的第一行必须是表头。"""
    path = (path or DEFAULT_DOCUMENT).expanduser()
    if not path.is_file():
        raise DemoError(f"表格文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        headers, raw_rows, failed_rows = _read_delimited(path, ",")
    elif suffix in {".txt", ".tsv"}:
        headers, raw_rows, failed_rows = _read_delimited(path, "\t")
    elif suffix in {".xlsx", ".xlsm"}:
        headers, raw_rows, failed_rows = _read_xlsx(path)
    elif suffix == ".xls":
        raise DemoError("旧版 XLS 不受 openpyxl 支持，请先转换为 XLSX")
    else:
        raise DemoError("仅支持 CSV、制表符 TXT/TSV 和 XLSX")

    # 与 table.py 一致，删除容易与内部索引冲突的保留字段。
    kept_indices = [index for index, header in enumerate(headers) if header.lower() not in RESERVED_COLUMNS]
    headers = [headers[index] for index in kept_indices]
    if not headers:
        raise DemoError("删除 id/index 等保留字段后没有可用列")
    columns = {header: [] for header in headers}
    for _, row in raw_rows:
        for output_index, source_index in enumerate(kept_indices):
            columns[headers[output_index]].append(row[source_index])
    column_types = {header: infer_column_type(columns[header]) for header in headers}

    records: list[TableRecord] = []
    for row_number, row in raw_rows:
        values = {
            headers[output_index]: _convert_value(row[source_index], column_types[headers[output_index]])
            for output_index, source_index in enumerate(kept_indices)
        }
        if any(value is not None for value in values.values()):
            records.append(TableRecord(row_number=row_number, values=values))
    return ParsedTable(path.name, headers, column_types, records, failed_rows)


def record_document(record: TableRecord) -> str:
    """按照 table.py 的格式，把参与索引的列拼成一行一个文档块。"""
    fields = []
    for name, value in record.values.items():
        role = COLUMN_ROLES.get(name, "both")
        if role not in INDEXING_ROLES or value is None:
            continue
        fields.append(f"- {name}: {value}")
    return "\n".join(fields)


def record_metadata(record: TableRecord, source: str) -> dict[str, str | int | float | bool]:
    """保存 metadata/both 列；Chroma metadata 只接受标量且不能保存 None。"""
    metadata: dict[str, str | int | float | bool] = {
        "source": source,
        "row_number": record.row_number,
        "parser": "table",
        "embedding_model": DEFAULT_MODEL,
    }
    for name, value in record.values.items():
        if COLUMN_ROLES.get(name, "both") in METADATA_ROLES and value is not None:
            metadata[name] = value
    return metadata


class QwenEmbeddingClient:
    """通过 OpenAI 兼容接口调用千问 text-embedding-v2。"""

    def __init__(self, api_key: str) -> None:
        if not api_key.strip():
            raise DemoError("API Key 不能为空")
        self.api_key = api_key
        self.url = DEFAULT_BASE_URL.rstrip("/") + "/embeddings"

    def _request(self, texts: list[str]) -> tuple[list[list[float]], int]:
        payload = json.dumps({"model": DEFAULT_MODEL, "input": texts, "encoding_format": "float"}).encode("utf-8")
        request = Request(self.url, data=payload, headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}, method="POST")
        for attempt in range(4):
            try:
                with urlopen(request, timeout=DEFAULT_TIMEOUT) as response:
                    body = json.load(response)
                items = sorted(body["data"], key=lambda item: item["index"])
                if [item["index"] for item in items] != list(range(len(texts))):
                    raise DemoError("向量接口返回的数量或顺序与输入不一致")
                vectors = [[float(value) for value in item["embedding"]] for item in items]
                if not vectors or not vectors[0] or any(len(vector) != len(vectors[0]) or not all(math.isfinite(value) for value in vector) for vector in vectors):
                    raise DemoError("向量为空、维度不一致或包含非法数值")
                return vectors, int((body.get("usage") or {}).get("total_tokens") or 0)
            except HTTPError as exc:
                detail = exc.read().decode("utf-8", errors="replace")[:500]
                if attempt == 3 or (exc.code not in {408, 429} and exc.code < 500):
                    raise DemoError(f"向量请求失败 HTTP {exc.code}: {detail}") from exc
            except (URLError, TimeoutError) as exc:
                if attempt == 3:
                    raise DemoError(f"无法连接向量接口：{exc}") from exc
            except (KeyError, TypeError, ValueError) as exc:
                raise DemoError(f"向量响应格式错误：{exc}") from exc
            time.sleep(min(2**attempt, 8))
        raise DemoError("向量请求失败")

    def encode(self, texts: list[str]) -> tuple[list[list[float]], int]:
        """按固定批次编码，确保各批向量维度相同。"""
        vectors: list[list[float]] = []
        usage = 0
        for start in range(0, len(texts), DEFAULT_BATCH_SIZE):
            batch, batch_usage = self._request(texts[start : start + DEFAULT_BATCH_SIZE])
            if vectors and len(batch[0]) != len(vectors[0]):
                raise DemoError("不同批次的向量维度不一致")
            vectors.extend(batch)
            usage += batch_usage
        return vectors, usage


class TableChromaStore:
    """持久化每行表格文本、外部生成的向量和字段元数据。"""

    def __init__(self) -> None:
        try:
            import chromadb
        except ImportError as exc:
            raise DemoError("请安装 chromadb：pip install -r requirements-table-demo.txt") from exc
        self.path = DEFAULT_CHROMA_PATH.resolve()
        self.client = chromadb.PersistentClient(path=str(self.path))
        self.collection = self.client.get_or_create_collection(
            name=DEFAULT_COLLECTION_NAME,
            embedding_function=None,
            metadata={"hnsw:space": "cosine", "embedding_model": DEFAULT_MODEL},
        )
        if (self.collection.metadata or {}).get("embedding_model") != DEFAULT_MODEL:
            raise DemoError("Chroma collection 的向量模型不匹配")

    def save(self, source: str, records: list[TableRecord], documents: list[str], vectors: list[list[float]]) -> list[str]:
        """同一文件重新运行时替换旧行，不影响 collection 中的其他文件。"""
        if not records or not (len(records) == len(documents) == len(vectors)):
            raise DemoError("记录、文本和向量数量必须一致且不能为空")
        dimension = len(vectors[0])
        if not dimension or any(len(vector) != dimension for vector in vectors):
            raise DemoError("向量维度为空或不一致")
        ids = [hashlib.sha256(f"{source}\0{record.row_number}\0{document}".encode("utf-8")).hexdigest() for record, document in zip(records, documents)]
        metadatas = [record_metadata(record, source) for record in records]
        self.collection.delete(where={"source": source})
        batch_size = self.client.get_max_batch_size()
        for start in range(0, len(ids), batch_size):
            end = start + batch_size
            self.collection.upsert(ids=ids[start:end], documents=documents[start:end], embeddings=vectors[start:end], metadatas=metadatas[start:end])
        return ids

    def query(self, vector: list[float]) -> list[dict]:
        count = min(DEFAULT_TOP_K, self.collection.count())
        if not count:
            return []
        result = self.collection.query(query_embeddings=[vector], n_results=count, include=["documents", "metadatas", "distances"])
        return [
            {"document": document, "metadata": metadata, "score": 1 - float(distance)}
            for document, metadata, distance in zip(result["documents"][0], result["metadatas"][0], result["distances"][0])
        ]


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """只保留运行 demo 必需的几个参数，其余配置使用代码顶部常量。"""
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("document", nargs="?", type=Path, default=DEFAULT_DOCUMENT, help=f"表格路径，默认：{DEFAULT_DOCUMENT.name}")
    parser.add_argument("--rows-only", action="store_true", help="只解析并打印每行 chunk，不调用 API 或写入 Chroma")
    parser.add_argument("--query", default=DEFAULT_QUERY)
    parser.add_argument("--api-key", default=os.getenv("DASHSCOPE_API_KEY") or os.getenv("QWEN_API_KEY"), help="默认读取 DASHSCOPE_API_KEY 或 QWEN_API_KEY")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """执行表格解析、每行分块、向量生成、Chroma 保存和查询。"""
    args = parse_args(argv)
    try:
        table = read_table(args.document)
        source = str(args.document.expanduser().resolve())
        usable = [(record, record_document(record)) for record in table.records]
        usable = [(record, document) for record, document in usable if document.strip()]
        if not usable:
            raise DemoError("表格没有可用于向量化的行，请检查字段或 COLUMN_ROLES")
    except (DemoError, OSError, csv.Error) as exc:
        print(f"Error: 表格解析失败：{exc}", file=sys.stderr)
        return 1

    records = [record for record, _ in usable]
    documents = [document for _, document in usable]
    print(f"document: {table.name}")
    print(f"columns: {len(table.headers)}, rows: {len(table.records)}, chunks: {len(documents)}")
    print("column types: " + ", ".join(f"{name}={table.column_types[name]}" for name in table.headers))
    if table.failed_rows:
        print("skipped malformed rows: " + ", ".join(map(str, table.failed_rows)))
    for index, (record, document) in enumerate(zip(records, documents), 1):
        print(f"\n--- row chunk {index}, source row={record.row_number} ---\n{document}")
    if args.rows_only:
        return 0

    api_key = args.api_key
    if not api_key and sys.stdin.isatty():
        api_key = getpass.getpass("Qwen API key: ")
    if not api_key:
        print("Error: 请设置 DASHSCOPE_API_KEY 或使用 --rows-only", file=sys.stderr)
        return 2
    try:
        client = QwenEmbeddingClient(api_key)
        store = TableChromaStore()
        vectors, document_usage = client.encode(documents)
        query_vectors, query_usage = client.encode([args.query])
        if len(vectors[0]) != len(query_vectors[0]):
            raise DemoError("文档向量和查询向量维度不一致")
        ids = store.save(source, records, documents, vectors)
        ranked = store.query(query_vectors[0])
    except Exception as exc:
        print(f"Error: 向量化或 Chroma 操作失败：{exc}", file=sys.stderr)
        return 1

    print("\n=== Chroma table retrieval ===")
    print(f"model: {DEFAULT_MODEL}, dimension: {len(query_vectors[0])}")
    print(f"provider token usage: document={document_usage}, query={query_usage}")
    print(f"chroma path: {store.path}\ncollection: {DEFAULT_COLLECTION_NAME}")
    print(f"stored rows: {len(ids)}, collection total: {store.collection.count()}\nquery: {args.query}")
    for rank, row in enumerate(ranked, 1):
        metadata = row["metadata"]
        preview = re.sub(r"\s+", " ", row["document"])
        print(f"{rank}. score={row['score']:.6f}, source row={metadata.get('row_number')}: {preview[:240]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
