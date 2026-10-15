#!/usr/bin/env python3
"""RAGFlow naive 文档分块与向量检索的独立运行示例。

本脚本不依赖 RAGFlow 的服务、数据库和配置中心。常见文本格式仅使用
Python 标准库解析，向量生成直接调用千问的 OpenAI 兼容接口，并将文档
分块、向量和元数据持久化到本地 Chroma 数据库。

运行示例：
    python naive_demo.py sample.md --chunks-only
    export DASHSCOPE_API_KEY='your-api-key'
    python naive_demo.py sample.md --query '向量检索是如何工作的？'
    python naive_demo.py --query '哪些因素会影响检索质量？'
    python naive_demo.py sample.md --chroma-path ./chroma_db
"""

from __future__ import annotations

import argparse
import csv
import getpass
import hashlib
import html
import json
import os
import re
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen


DEFAULT_BASE_URL = (
    "https://ws-leyp0n1kyjsllxss.cn-beijing.maas.aliyuncs.com/"
    "compatible-mode/v1"
)
DEFAULT_MODEL = "text-embedding-v2"
# Chroma 默认保存在 demo 文件同级目录，避免因运行时工作目录不同而找不到数据库。
DEFAULT_CHROMA_PATH = Path(__file__).resolve().parent / "chroma_db"
DEFAULT_COLLECTION_NAME = "naive-demo"
# 与 rag/app/naive.py 的默认分隔符保持一致，包含换行和中英文句末符号。
DEFAULT_DELIMITER = "\n!?;。；！？"
# 未传入文档路径时使用这段内置文本，方便直接运行脚本验证完整流程。
DEFAULT_SAMPLE = """# 企业知识库检索

RAG 系统通常先将文档解析并切分为较小的文本块，然后计算每个文本块的向量。
合理的块大小能够保留上下文，同时减少无关内容对相似度的干扰。

用户提问时，系统使用相同的向量模型生成查询向量，并通过余弦相似度召回相关文本块。
召回数量、文本切分策略、向量模型和重排序模型都会影响最终回答的准确率。

结构化数据库索引适合精确查询，但不能直接替代语义向量检索。
"""


class DemoError(RuntimeError):
    """文档解析或向量调用失败时抛出的统一异常。"""


def token_count(text: str) -> int:
    """在不引入 RAGFlow tiktoken 依赖的情况下估算 Token 数。

    中文字符逐字计数，连续英文、数字和下划线按一个 Token 计算，其他非空白
    符号单独计数。该数值只用于控制分块大小；向量接口的实际用量仍以服务端
    返回的 usage 为准。
    """
    return len(re.findall(r"[\u3400-\u9fff]|[A-Za-z0-9_]+|[^\s]", text or ""))


def _custom_delimiters(delimiter: str) -> list[str]:
    """提取反引号包裹的自定义分隔符，例如 `` `##` ``。"""
    return sorted(set(re.findall(r"`([^`]+)`", delimiter)), key=len, reverse=True)


def _normal_delimiter_pattern(delimiter: str) -> str:
    """将普通分隔字符转换成可用于 re.split 的正则表达式。"""
    # 自定义分隔符有独立语义，构建普通字符正则时需要先移除。
    plain = re.sub(r"`[^`]+`", "", delimiter)
    return "|".join(re.escape(char) for char in dict.fromkeys(plain) if char)


def _split_at_delimiters(text: str, pattern: str) -> list[str]:
    """按句子分隔符拆分文本，并把分隔符保留在前一个文本片段末尾。"""
    if not pattern:
        return [text]
    # 使用捕获组保留分隔符，避免切分后丢失原文中的标点。
    pieces = re.split(f"({pattern})", text, flags=re.DOTALL)
    result: list[str] = []
    for piece in pieces:
        if not piece:
            continue
        if re.fullmatch(pattern, piece):
            if result:
                result[-1] += piece
            continue
        result.append(piece)
    return result


def naive_merge(
    sections: str | Iterable[str],
    chunk_token_num: int = 128,
    delimiter: str = DEFAULT_DELIMITER,
    overlapped_percent: float = 0,
) -> list[str]:
    """使用 ``rag/app/naive.py`` 的核心策略合并文本段落。

    超长段落先按配置的分隔符拆开，小段落则持续合并到目标大小。使用反引号
    包裹的自定义分隔符（例如 `` `##` ``）时，每个分隔片段直接成为独立块，
    这一行为与 RAGFlow 的 naive_merge 保持一致。
    """
    # 统一将输入转换为非空字符串列表，兼容单段文本和多段文本两种调用方式。
    if isinstance(sections, str):
        values = [sections]
    else:
        values = [section for section in sections if section and section.strip()]
    if not values:
        return []
    if chunk_token_num < 1:
        return ["\n".join(values)]
    # 限制重叠比例，避免异常参数导致新分块几乎完全复制上一块。
    overlap = min(max(float(overlapped_percent), 0.0), 90.0)

    # 自定义分隔符优先级最高，不再按照 chunk_token_num 对片段进行二次合并。
    custom = _custom_delimiters(delimiter)
    if custom:
        pattern = "|".join(re.escape(item) for item in custom)
        return [
            piece.strip()
            for section in values
            for piece in re.split(pattern, section, flags=re.DOTALL)
            if piece.strip()
        ]

    # 普通模式下，先把超过目标大小的段落拆成句子级单元。
    delimiter_pattern = _normal_delimiter_pattern(delimiter)
    units: list[str] = []
    for section in values:
        if token_count(section) < chunk_token_num or not delimiter_pattern:
            units.append(section)
        else:
            units.extend(_split_at_delimiters(section, delimiter_pattern))

    # 再按顺序把句子单元合并为最终分块。
    chunks: list[str] = []
    current = ""
    current_tokens = 0
    threshold = chunk_token_num * (100.0 - overlap) / 100.0

    for unit in units:
        text = unit if not current else f"\n{unit}"
        unit_tokens = token_count(text)
        if current and current_tokens > threshold:
            # 当前块达到阈值后封口，并从尾部截取指定比例作为下一块上下文。
            chunks.append(current.strip())
            overlap_text = ""
            if overlap:
                overlap_chars = int(len(current) * overlap / 100.0)
                overlap_text = current[-overlap_chars:] if overlap_chars else ""
            current = f"{overlap_text}{text}"
            current_tokens = token_count(current)
        else:
            current += text
            current_tokens += unit_tokens

    if current.strip():
        chunks.append(current.strip())
    return chunks


class _TextHTMLParser(HTMLParser):
    """将 HTML 转成纯文本，同时尽量保留块级标签形成的段落边界。"""

    # 遇到这些标签时插入换行，避免相邻段落被直接拼接。
    BLOCK_TAGS = {
        "article",
        "blockquote",
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "p",
        "section",
        "table",
        "tr",
    }
    # 脚本和样式内容对语义检索没有帮助，解析时直接忽略。
    SKIP_TAGS = {"script", "style"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.SKIP_TAGS:
            self.skip_depth += 1
        elif not self.skip_depth and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP_TAGS and self.skip_depth:
            self.skip_depth -= 1
        elif not self.skip_depth and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip_depth:
            self.parts.append(data)

    def text(self) -> str:
        value = html.unescape("".join(self.parts))
        value = re.sub(r"[ \t]+", " ", value)
        return re.sub(r"\n\s*\n+", "\n\n", value).strip()


def _decode(data: bytes) -> str:
    """按常见中英文编码顺序尝试解码文件内容。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _json_sections(text: str, json_lines: bool) -> list[str]:
    """把 JSON 或 JSON Lines 转换为保持对象边界的文本段落。"""
    if json_lines:
        objects = [json.loads(line) for line in text.splitlines() if line.strip()]
    else:
        value = json.loads(text)
        objects = value if isinstance(value, list) else [value]
    return [json.dumps(item, ensure_ascii=False, indent=2) for item in objects]


def _docx_text(path: Path) -> str:
    """提取 DOCX 的段落和表格文本；该格式需要 python-docx。"""
    try:
        from docx import Document
    except ImportError as exc:
        raise DemoError("DOCX parsing requires: pip install python-docx") from exc
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _pdf_text(path: Path) -> str:
    """逐页提取 PDF 文本；扫描件仍需额外 OCR，本示例不处理 OCR。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DemoError("PDF parsing requires: pip install pypdf") from exc
    return "\n\n".join((page.extract_text() or "") for page in PdfReader(path).pages)


def read_sections(path: Path | None) -> tuple[str, list[str]]:
    """根据文件扩展名读取文档，并返回文档名与有序文本段落。"""
    # 不传文件时使用内置示例，使脚本开箱即可测试。
    if path is None:
        return "built-in-sample.md", [DEFAULT_SAMPLE]
    if not path.is_file():
        raise DemoError(f"Document does not exist: {path}")

    # PDF 和 DOCX 需要可选第三方依赖，因此在通用文本解码前单独处理。
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return path.name, [_docx_text(path)]
    if suffix == ".pdf":
        return path.name, [_pdf_text(path)]

    # 其他格式先读取为文本，再按照数据结构做进一步拆分。
    text = _decode(path.read_bytes())
    if suffix in {".htm", ".html"}:
        parser = _TextHTMLParser()
        parser.feed(text)
        return path.name, [parser.text()]
    if suffix == ".json":
        return path.name, _json_sections(text, json_lines=False)
    if suffix in {".jsonl", ".ldjson"}:
        return path.name, _json_sections(text, json_lines=True)
    if suffix == ".csv":
        rows = csv.reader(text.splitlines())
        return path.name, ["\t".join(row) for row in rows]

    supported = {
        ".txt",
        ".md",
        ".markdown",
        ".mdx",
        ".py",
        ".js",
        ".ts",
        ".java",
        ".c",
        ".cpp",
        ".h",
        ".go",
        ".php",
        ".sh",
        ".sql",
        ".cs",
        ".kt",
    }
    if suffix not in supported:
        raise DemoError(f"Unsupported file type: {suffix or '<none>'}")
    return path.name, [text]


def embeddings_url(base_url: str) -> str:
    """补齐 OpenAI 兼容接口的 /embeddings 路径，避免重复拼接。"""
    normalized = base_url.rstrip("/")
    return normalized if normalized.endswith("/embeddings") else f"{normalized}/embeddings"


class QwenEmbeddingClient:
    """面向千问 text-embedding-v2 的轻量 OpenAI 兼容客户端。"""

    def __init__(
        self,
        api_key: str,
        base_url: str = DEFAULT_BASE_URL,
        model: str = DEFAULT_MODEL,
        batch_size: int = 4,
        timeout: float = 60.0,
        max_retries: int = 3,
    ) -> None:
        if not api_key.strip():
            raise ValueError("API key cannot be empty")
        self.api_key = api_key
        self.url = embeddings_url(base_url)
        self.model = model
        self.batch_size = max(1, batch_size)
        self.timeout = timeout
        self.max_retries = max(0, max_retries)

    def _request(self, texts: list[str]) -> tuple[list[list[float]], int]:
        """发送一批文本，并返回与输入顺序一致的向量及 Token 用量。"""
        # encoding_format=float 确保服务端直接返回浮点数组，而不是 Base64 数据。
        payload = json.dumps(
            {"model": self.model, "input": texts, "encoding_format": "float"}
        ).encode("utf-8")
        request = Request(
            self.url,
            data=payload,
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )

        # 对限流、超时和服务端错误进行指数退避重试；普通 4XX 不重试。
        last_error: Exception | None = None
        for attempt in range(self.max_retries + 1):
            retryable = False
            try:
                with urlopen(request, timeout=self.timeout) as response:
                    body = json.loads(response.read().decode("utf-8"))
                items = body.get("data")
                if not isinstance(items, list):
                    raise DemoError(f"Invalid embedding response: {str(body)[:500]}")
                # 某些兼容接口不保证返回顺序，按 index 排序后再与原文对应。
                items.sort(key=lambda item: item.get("index", 0))
                vectors = [[float(value) for value in item["embedding"]] for item in items]
                if len(vectors) != len(texts):
                    raise DemoError(
                        f"Service returned {len(vectors)} vectors for {len(texts)} texts"
                    )
                usage = body.get("usage") or {}
                return vectors, int(usage.get("total_tokens") or 0)
            except HTTPError as exc:
                detail = exc.read().decode("utf-8", errors="replace")
                last_error = DemoError(f"HTTP {exc.code}: {detail[:500]}")
                retryable = exc.code in {408, 429} or exc.code >= 500
            except (URLError, TimeoutError) as exc:
                last_error = DemoError(str(exc))
                retryable = True
            except (KeyError, TypeError, ValueError, json.JSONDecodeError, DemoError) as exc:
                last_error = exc if isinstance(exc, DemoError) else DemoError(str(exc))

            if not retryable or attempt == self.max_retries:
                break
            time.sleep(min(2**attempt, 8))
        raise DemoError(f"Embedding request failed: {last_error}") from last_error

    def encode(self, texts: Iterable[str]) -> tuple[list[list[float]], int]:
        """按 batch_size 批量编码文档文本，并累计服务端 Token 用量。"""
        values = list(texts)
        vectors: list[list[float]] = []
        total_tokens = 0
        for start in range(0, len(values), self.batch_size):
            batch_vectors, batch_tokens = self._request(
                values[start : start + self.batch_size]
            )
            vectors.extend(batch_vectors)
            total_tokens += batch_tokens
        return vectors, total_tokens

    def encode_queries(self, text: str) -> tuple[list[float], int]:
        """编码单条查询文本，返回一维向量及 Token 用量。"""
        vectors, total_tokens = self.encode([text])
        return vectors[0], total_tokens


class ChromaVectorStore:
    """使用 Chroma PersistentClient 保存并检索外部生成的向量。"""

    def __init__(self, path: Path, collection_name: str) -> None:
        try:
            import chromadb
        except ImportError as exc:
            raise DemoError(
                "缺少 chromadb，请执行：pip install -r "
                "rag/app/test/requirements-naive-demo.txt"
            ) from exc

        # PersistentClient 会在指定目录中保存数据库，并在下次启动时重新加载。
        self.path = path.expanduser().resolve()
        self.path.mkdir(parents=True, exist_ok=True)
        self.collection_name = collection_name
        self.client = chromadb.PersistentClient(path=str(self.path))
        self.collection = self.client.get_or_create_collection(
            name=collection_name,
            # Chroma 0.6.x 和新版本都支持通过 metadata 配置 HNSW 距离空间。
            metadata={"hnsw:space": "cosine"},
        )

    @staticmethod
    def _chunk_id(source: str, chunk_index: int, chunk: str) -> str:
        """生成稳定 ID，使同一文档重复运行时可以更新原有分块。"""
        value = f"{source}\0{chunk_index}\0{chunk}".encode("utf-8")
        return hashlib.sha256(value).hexdigest()

    def save(
        self,
        source: str,
        chunks: list[str],
        vectors: list[list[float]],
        model: str,
    ) -> list[str]:
        """删除该文档的旧分块，再将最新文本、向量和元数据写入 Chroma。"""
        if len(chunks) != len(vectors):
            raise DemoError(
                f"分块数量和向量数量不一致：{len(chunks)} != {len(vectors)}"
            )

        # 先按 source 清理旧数据，避免文档变短后遗留已经不存在的旧分块。
        self.collection.delete(where={"source": source})
        ids = [
            self._chunk_id(source, index, chunk)
            for index, chunk in enumerate(chunks, start=1)
        ]
        metadatas = [
            {
                "source": source,
                "chunk_index": index,
                "estimated_tokens": token_count(chunk),
                "embedding_model": model,
            }
            for index, chunk in enumerate(chunks, start=1)
        ]
        self.collection.upsert(
            ids=ids,
            documents=chunks,
            embeddings=vectors,
            metadatas=metadatas,
        )
        return ids

    def query(self, query_vector: list[float], top_k: int) -> list[dict]:
        """使用查询向量检索整个 collection，并返回整理后的结果。"""
        result_count = min(max(0, top_k), self.collection.count())
        if result_count == 0:
            return []

        result = self.collection.query(
            query_embeddings=[query_vector],
            n_results=result_count,
            include=["documents", "metadatas", "distances"],
        )
        ids = (result.get("ids") or [[]])[0]
        documents = (result.get("documents") or [[]])[0]
        metadatas = (result.get("metadatas") or [[]])[0]
        distances = (result.get("distances") or [[]])[0]

        rows = []
        for record_id, document, metadata, distance in zip(
            ids, documents, metadatas, distances
        ):
            # cosine 距离越小越相似，转换成更直观的相似度分数：1 - distance。
            rows.append(
                {
                    "id": record_id,
                    "document": document or "",
                    "metadata": metadata or {},
                    "distance": float(distance),
                    "score": 1.0 - float(distance),
                }
            )
        return rows


def parse_args() -> argparse.Namespace:
    """定义命令行参数。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document", nargs="?", type=Path, help="Document to parse")
    parser.add_argument("--query", default="什么因素会影响知识库检索准确率？")
    parser.add_argument("--chunk-size", type=int, default=128)
    parser.add_argument("--overlap", type=float, default=10.0, help="Overlap percentage")
    parser.add_argument("--delimiter", default=DEFAULT_DELIMITER)
    parser.add_argument("--top-k", type=int, default=3)
    parser.add_argument("--chunks-only", action="store_true", help="Skip embedding calls")
    parser.add_argument(
        "--api-key",
        default=os.getenv("DASHSCOPE_API_KEY") or os.getenv("QWEN_API_KEY"),
        help="Defaults to DASHSCOPE_API_KEY or QWEN_API_KEY",
    )
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--batch-size", type=int, default=4)
    parser.add_argument("--timeout", type=float, default=60.0)
    parser.add_argument(
        "--chroma-path",
        type=Path,
        default=DEFAULT_CHROMA_PATH,
        help=f"Chroma 持久化目录，默认：{DEFAULT_CHROMA_PATH}",
    )
    parser.add_argument(
        "--collection-name",
        default=DEFAULT_COLLECTION_NAME,
        help=f"Chroma collection 名称，默认：{DEFAULT_COLLECTION_NAME}",
    )
    return parser.parse_args()


def main() -> int:
    """执行文档解析、naive 分块、向量生成和 Top-K 检索。"""
    args = parse_args()
    try:
        # 第一阶段：将输入文件解析成 sections，再执行 naive_merge。
        document_name, sections = read_sections(args.document)
        chunks = naive_merge(
            sections,
            chunk_token_num=args.chunk_size,
            delimiter=args.delimiter,
            overlapped_percent=args.overlap,
        )
    except (DemoError, OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print(f"document: {document_name}")
    print(f"sections: {len(sections)}, chunks: {len(chunks)}")
    for index, chunk in enumerate(chunks, start=1):
        print(f"\n--- chunk {index} ({token_count(chunk)} estimated tokens) ---")
        print(chunk)

    # --chunks-only 适合离线调试分块，无需 API key，也不会发起网络请求。
    if args.chunks_only:
        return 0
    if not chunks:
        print("Error: no text chunks were produced", file=sys.stderr)
        return 1

    # 优先读取命令行或环境变量；交互终端中缺失时使用隐藏输入，避免明文回显。
    api_key = args.api_key
    if not api_key and sys.stdin.isatty():
        api_key = getpass.getpass("Qwen API key: ")
    if not api_key:
        print(
            "Error: set DASHSCOPE_API_KEY, pass --api-key, or use --chunks-only",
            file=sys.stderr,
        )
        return 2

    # 第二阶段：创建千问向量客户端，对所有分块和用户查询分别编码。
    client = QwenEmbeddingClient(
        api_key=api_key,
        base_url=args.base_url,
        model=args.model,
        batch_size=args.batch_size,
        timeout=args.timeout,
    )
    try:
        chunk_vectors, document_tokens = client.encode(chunks)
        query_vector, query_tokens = client.encode_queries(args.query)
        # 第三阶段：将文档块和向量持久化到 Chroma，再由 Chroma 完成 Top-K 检索。
        source = str(args.document.expanduser().resolve()) if args.document else document_name
        vector_store = ChromaVectorStore(args.chroma_path, args.collection_name)
        stored_ids = vector_store.save(source, chunks, chunk_vectors, args.model)
        ranked = vector_store.query(query_vector, args.top_k)
    except (DemoError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    print("\n=== vector retrieval ===")
    print(f"model: {args.model}")
    print(f"endpoint: {client.url}")
    print(f"embedding dimension: {len(query_vector)}")
    print(f"provider token usage: document={document_tokens}, query={query_tokens}")
    print(f"chroma path: {vector_store.path}")
    print(f"collection: {vector_store.collection_name}")
    print(f"stored chunks: {len(stored_ids)}, collection total: {vector_store.collection.count()}")
    print(f"query: {args.query}")
    for rank, row in enumerate(ranked, start=1):
        metadata = row["metadata"]
        preview = re.sub(r"\s+", " ", row["document"]).strip()
        print(
            f"{rank}. score={row['score']:.6f}, "
            f"source={metadata.get('source')}, "
            f"chunk={metadata.get('chunk_index')}: {preview[:240]}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
